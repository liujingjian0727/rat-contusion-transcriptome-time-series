#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import argparse
import csv
import math
import re
import statistics
import sys
from collections import Counter, defaultdict
from pathlib import Path


EXTS = {".tsv", ".txt", ".csv", ".xlsx", ".xlsm"}

MISS = {"", "na", "n/a", "null", "none", "."}
BADNAN = {"nan"}
BADINF = {"inf", "+inf", "-inf", "infinity", "+infinity", "-infinity"}


# ============================================================
# Frozen manuscript expectations
# ============================================================

S1_COUNTS = {
    "Baseline": 7,
    "0h": 9,
    "1h": 8,
    "2h": 7,
    "6h": 6,
    "10h": 6,
    "14h": 7,
    "18h": 6,
    "36h": 7,
    "60h": 7,
    "72h": 7,
}

S5_COUNTS = {
    "Baseline": 6,
    "4h": 7,
    "8h": 8,
    "12h": 8,
    "16h": 8,
    "20h": 7,
    "24h": 8,
    "48h": 8,
}

PRIMARY = [
    "0h", "1h", "2h", "6h", "10h",
    "14h", "18h", "36h", "60h", "72h"
]

HIST = [
    "4h", "8h", "12h", "16h",
    "20h", "24h", "48h"
]

DEG2026 = {
    "0h": (760, 901, 1661),
    "1h": (1344, 2389, 3733),
    "2h": (1088, 1237, 2325),
    "6h": (1158, 903, 2061),
    "10h": (477, 977, 1454),
    "14h": (1169, 1435, 2604),
    "18h": (1583, 1892, 3475),
    "36h": (1207, 1422, 2629),
    "60h": (1610, 1579, 3189),
    "72h": (2230, 2224, 4454),
}

DEG2021 = {
    "4h": (2363, 1675, 4038),
    "8h": (2516, 2314, 4830),
    "12h": (1729, 963, 2692),
    "16h": (1549, 1368, 2917),
    "20h": (978, 538, 1516),
    "24h": (1781, 1847, 3628),
    "48h": (2173, 1403, 3576),
}

PEAK476 = {
    "0h": 28,
    "1h": 36,
    "2h": 38,
    "6h": 25,
    "10h": 6,
    "14h": 8,
    "18h": 24,
    "36h": 8,
    "60h": 13,
    "72h": 290,
}


# ============================================================
# Column aliases
# ============================================================

ALIASES = {
    "sample_id": [
        "sample_id", "sampleid", "sample",
        "sample_name", "samplename"
    ],

    "sampling_stage": [
        "sampling_stage", "samplingstage",
        "stage", "group", "timepoint",
        "time_point", "time_factor"
    ],

    "time_h": [
        "time_h", "timeh", "time_hour",
        "time_hours", "hour", "hours", "time"
    ],

    "condition": [
        "condition", "injury_condition",
        "status", "experimental_condition"
    ],

    "cohort": [
        "cohort", "dataset",
        "cohort_id", "dataset_id"
    ],

    "gene_id": [
        "gene_id", "geneid",
        "ensembl_gene_id",
        "ensemblgeneid", "gene"
    ],

    "contrast": [
        "contrast", "comparison", "de_contrast"
    ],

    "n_upregulated": [
        "n_upregulated", "upregulated",
        "upregulated_genes", "up_genes",
        "up_deg", "up_degs", "n_up", "up"
    ],

    "n_downregulated": [
        "n_downregulated", "downregulated",
        "downregulated_genes", "down_genes",
        "down_deg", "down_degs", "n_down", "down"
    ],

    "n_total_deg": [
        "n_total_deg", "total_deg",
        "total_degs", "n_deg", "n_degs",
        "deg_total", "total"
    ],

    "n_genes_tested": [
        "n_genes_tested", "genes_tested",
        "n_tested", "tested_genes"
    ],

    "n_samples_baseline": [
        "n_samples_baseline", "baseline_n",
        "n_baseline", "control_n"
    ],

    "spline_lrt_padj": [
        "spline_lrt_padj", "lrt_padj",
        "spline_padj", "padj", "fdr"
    ],

    "lrt_dynamic_fdr005": [
        "lrt_dynamic_fdr005",
        "lrt_dynamic_fdr_005",
        "time_associated_fdr05"
    ],

    "tau085_spm050_injury10": [
        "tau085_and_spm050_injury10",
        "tau085_spm050_injury10",
        "tau_spm_injury10_pass"
    ],

    "mean_vst_temporal_range": [
        "mean_vst_temporal_range",
        "vst_temporal_range",
        "temporal_range"
    ],

    "tau_injury10": [
        "tau_injury10", "tau",
        "injury_tau", "tau_injury"
    ],

    "peak_spm_injury10": [
        "peak_spm_injury10",
        "peak_spm", "spm", "max_spm"
    ],

    "higher_amplitude_dynamic": [
        "higher_amplitude_dynamic",
        "high_amplitude_dynamic",
        "amplitude_pass",
        "high_confidence_dynamic"
    ],

    "temporal_specificity_pass": [
        "temporal_specificity_pass",
        "specificity_pass",
        "tau_spm_pass",
        "temporal_specific"
    ],

    "observed_peak_stage": [
        "observed_peak_stage",
        "peak_stage",
        "peak_timepoint",
        "peak_time",
        "peak_group",
        "peak_stage_injury10",
        "peak_stage_label_injury10"
    ],

    "primary_stage": [
        "primary_stage",
        "stage_2026",
        "cohort2026_stage",
        "a_stage",
        "a_timepoint",
        "primary_timepoint"
    ],

    "historical_stage": [
        "historical_stage",
        "stage_2021",
        "cohort2021_stage",
        "r_stage",
        "r_timepoint",
        "historical_timepoint"
    ],

    "primary_time_h": [
        "primary_time_h",
        "time_2026_h",
        "a_time_h",
        "a_time",
        "primary_time"
    ],

    "historical_time_h": [
        "historical_time_h",
        "time_2021_h",
        "r_time_h",
        "r_time",
        "historical_time"
    ],

    "absolute_time_gap_h": [
        "absolute_time_gap_h",
        "abs_time_gap_h",
        "time_gap_h"
    ],

    "nearest_time_pair": [
        "nearest_time_pair",
        "nearest_pair",
        "chronologically_nearest"
    ],

    "n_common_genes": [
        "n_common_genes",
        "common_genes",
        "n_shared_genes_universe"
    ],

    "spearman_rho": [
        "spearman_rho",
        "spearman_cor",
        "spearman_corr",
        "spearman_correlation",
        "spearman_r",
        "rho"
    ],

    "pearson_r": [
        "pearson_r",
        "pearson_cor",
        "pearson_corr",
        "pearson_correlation"
    ],

    "shared_deg_n": [
        "shared_deg_n",
        "n_shared_deg",
        "shared_degs",
        "shared_deg"
    ],

    "shared_deg_same_direction_n": [
        "shared_deg_same_direction_n",
        "same_direction_n",
        "n_same_direction"
    ],

    "direction_concordance_percent": [
        "direction_concordance_percent",
        "direction_concordance",
        "same_direction_percent",
        "direction_concordance_pct"
    ],

    "scenario_id": [
        "scenario_id",
        "scenario",
        "analysis_scenario",
        "scenario_name",
        "scenario_or_removed_sample",
        "removed_sample",
        "scenario_or_sample"
    ],

    "n_excluded": [
        "n_excluded",
        "excluded_n",
        "number_excluded"
    ],

    "n_samples_remaining": [
        "n_samples_remaining",
        "n_remaining",
        "remaining_n",
        "sample_n_remaining",
        "n_samples"
    ],

    "permanova_r2": [
        "permanova_r2",
        "permanova_rsq",
        "r2_permanova"
    ],

    "permanova_p": [
        "permanova_p",
        "permanova_pvalue",
        "permanova_p_value"
    ],

    "permdisp_p": [
        "permdisp_p",
        "permdisp_pvalue",
        "permdisp_p_value"
    ],

    "matrix_correlation_vs_primary": [
        "matrix_correlation_vs_primary",
        "matrix_cor_vs_primary",
        "matrix_correlation",
        "correlation_vs_primary",
        "cross_cohort_matrix_correlation"
    ],

    "matrix_mean_absolute_difference": [
        "matrix_mean_absolute_difference",
        "matrix_mae",
        "mean_absolute_difference",
        "mae",
        "cross_cohort_matrix_mae"
    ],
}


# ============================================================
# Helper functions
# ============================================================

def nh(x):
    return re.sub(
        r"[^a-z0-9]+",
        "",
        str(x).strip().lower()
    )


def missing(x):
    if x is None:
        return True

    if isinstance(x, float) and math.isnan(x):
        return False

    return str(x).strip().lower() in MISS


def badnum(x):
    if isinstance(x, float):
        if math.isnan(x):
            return "NaN"
        if math.isinf(x):
            return "Inf"
        return None

    s = str(x).strip().lower()

    if s in BADNAN:
        return "NaN"

    if s in BADINF:
        return "Inf"

    return None


def num(x):
    if missing(x) or badnum(x):
        return None

    s = str(x).strip().replace(",", "")

    if s.endswith("%"):
        s = s[:-1].strip()

    try:
        v = float(s)
    except Exception:
        return None

    if not math.isfinite(v):
        return None

    return v


def integer(x):
    v = num(x)

    if v is None:
        return None

    if abs(v - round(v)) < 1e-8:
        return int(round(v))

    return None


def boolean(x):
    if missing(x):
        return None

    if isinstance(x, bool):
        return x

    s = str(x).strip().lower()

    if s in {
        "true", "t", "yes", "y",
        "1", "pass", "passed"
    }:
        return True

    if s in {
        "false", "f", "no", "n",
        "0", "fail", "failed"
    }:
        return False

    return None


def stage(x):
    if missing(x):
        return None

    s = str(x).strip()

    z = re.sub(
        r"[\s_\-]+",
        "",
        s.lower()
    )

    if z in {
        "baseline",
        "control",
        "uninjured",
        "normal",
        "normalcontrol"
    }:
        return "Baseline"

    m = re.fullmatch(
        r"[ra]?(\d+(?:\.\d+)?)h",
        z
    )

    if not m:
        m = re.fullmatch(
            r"(\d+(?:\.\d+)?)",
            z
        )

    if m:
        v = float(m.group(1))

        if v.is_integer():
            return f"{int(v)}h"

        return f"{v:g}h"

    return s


def stage_time(s):
    if not s or s == "Baseline":
        return None

    m = re.fullmatch(
        r"(\d+(?:\.\d+)?)h",
        s
    )

    if not m:
        return None

    return float(m.group(1))


def cond(x):
    if missing(x):
        return None

    z = re.sub(
        r"[\s_\-]+",
        "",
        str(x).lower()
    )

    if z in {
        "uninjured",
        "control",
        "normal",
        "baseline"
    }:
        return "Uninjured"

    if z in {
        "injury",
        "injured",
        "contusion",
        "treatment",
        "treated"
    }:
        return "Injury"

    return str(x).strip()


# ============================================================
# Data structures
# ============================================================

class Table:

    def __init__(
        self,
        key,
        path,
        headers,
        rows
    ):
        self.key = key
        self.path = path
        self.headers = headers
        self.rows = rows

    @property
    def n(self):
        return len(self.rows)


class Audit:

    def __init__(self, strict=False):
        self.strict = strict
        self.rows = []

    def add(
        self,
        table,
        check_id,
        status,
        message,
        observed="",
        expected="",
        path=""
    ):

        status = status.upper()

        if self.strict and status == "WARN":
            status = "FAIL"
            message = "[strict] " + message

        self.rows.append([
            table,
            check_id,
            status,
            message,
            str(observed),
            str(expected),
            str(path)
        ])

    def failed(self):
        return any(
            r[2] == "FAIL"
            for r in self.rows
        )


# ============================================================
# Column resolution
# ============================================================

def col(t, canonical):

    m = {
        nh(h): h
        for h in t.headers
    }

    for a in [canonical] + ALIASES.get(canonical, []):

        if nh(a) in m:
            return m[nh(a)]

    return None


def require(a, t, names):

    d = {
        x: col(t, x)
        for x in names
    }

    miss = [
        x for x, v in d.items()
        if not v
    ]

    if miss:

        a.add(
            t.key,
            "required_columns",
            "FAIL",
            "Required columns could not be resolved",
            ";".join(t.headers),
            ";".join(names),
            t.path
        )

    else:

        a.add(
            t.key,
            "required_columns",
            "PASS",
            "Required columns resolved",
            "; ".join(
                f"{k}={v}"
                for k, v in d.items()
            ),
            "",
            t.path
        )

    return d


def optional(a, t, name):

    c = col(t, name)

    if not c:

        a.add(
            t.key,
            "optional_" + name,
            "WARN",
            f"Recommended column '{name}' not found",
            path=t.path
        )

    return c


# ============================================================
# Input
# ============================================================

def read_table(key, p):

    if p.suffix.lower() in {
        ".xlsx",
        ".xlsm"
    }:

        try:
            from openpyxl import load_workbook

        except ImportError:

            raise RuntimeError(
                "openpyxl required for XLSX/XLSM"
            )

        wb = load_workbook(
            p,
            read_only=True,
            data_only=True
        )

        vals = None

        for ws in wb.worksheets:

            x = [
                list(r)
                for r in ws.iter_rows(
                    values_only=True
                )
                if any(
                    v is not None
                    and str(v).strip()
                    for v in r
                )
            ]

            if x:
                vals = x
                break

        if not vals:
            raise ValueError(
                "No non-empty worksheet"
            )

        headers = [
            ""
            if v is None
            else str(v).strip()
            for v in vals[0]
        ]

        rows = []

        for r in vals[1:]:

            r = (
                r + [None] * len(headers)
            )[:len(headers)]

            rows.append(
                dict(
                    zip(
                        headers,
                        r
                    )
                )
            )

    else:

        with p.open(
            "r",
            encoding="utf-8-sig",
            newline=""
        ) as f:

            sample = f.read(8192)
            f.seek(0)

            if (
                p.suffix.lower() == ".tsv"
                or "\t" in sample
            ):
                delim = "\t"
            else:
                delim = ","

            rr = [
                r
                for r in csv.reader(
                    f,
                    delimiter=delim
                )
                if any(
                    str(v).strip()
                    for v in r
                )
            ]

        if not rr:
            raise ValueError(
                "Empty file"
            )

        headers = [
            str(v).strip()
            for v in rr[0]
        ]

        rows = []

        for r in rr[1:]:

            r = (
                r + [""] * len(headers)
            )[:len(headers)]

            rows.append(
                dict(
                    zip(
                        headers,
                        r
                    )
                )
            )

    return Table(
        key,
        p,
        headers,
        rows
    )


# ============================================================
# Auto-discovery
# ============================================================

def discover(root, n):

    pat = re.compile(
        rf"(?:^|[^a-z0-9])s0?{n}(?:[^0-9]|$)",
        re.I
    )

    cand = []

    for p in root.rglob("*"):

        if not p.is_file():
            continue

        if p.suffix.lower() not in EXTS:
            continue

        if not pat.search(p.stem):
            continue

        if any(
            x in p.name.lower()
            for x in [
                "index",
                "manifest",
                "validation"
            ]
        ):
            continue

        if p.suffix.lower() == ".tsv":
            score = 30

        elif p.suffix.lower() == ".csv":
            score = 20

        else:
            score = 10

        if (
            f"supplementarytables{n}"
            in nh(p.stem)
        ):
            score += 100

        if (
            p.parent.resolve()
            == root.resolve()
        ):
            score += 10

        cand.append(
            (
                score,
                p
            )
        )

    if not cand:
        return None, []

    cand = sorted(
        cand,
        key=lambda x: (
            -x[0],
            str(x[1])
        )
    )

    top = cand[0][0]

    best = [
        p
        for s, p in cand
        if s == top
    ]

    if len(best) == 1:

        return (
            best[0],
            [
                p
                for _, p in cand
            ]
        )

    return None, best


# ============================================================
# Generic validation
# ============================================================

def generic(a, t):

    a.add(
        t.key,
        "nonempty",
        "PASS" if t.n else "FAIL",
        "Table has data rows"
        if t.n
        else "No data rows",
        t.n,
        "",
        t.path
    )

    nn = [
        nh(h)
        for h in t.headers
    ]

    dup = [
        k
        for k, v in Counter(nn).items()
        if k and v > 1
    ]

    a.add(
        t.key,
        "unique_headers",
        "FAIL" if dup else "PASS",
        "Duplicate normalized headers"
        if dup
        else "Headers unique",
        dup,
        "",
        t.path
    )

    nan = 0
    inf = 0

    for r in t.rows:

        for v in r.values():

            b = badnum(v)

            nan += b == "NaN"
            inf += b == "Inf"

    a.add(
        t.key,
        "nan_inf",
        "FAIL"
        if nan or inf
        else "PASS",
        "NaN/Inf detected"
        if nan or inf
        else "No NaN/Inf detected",
        f"NaN={nan};Inf={inf}",
        "Use NA; no Inf",
        t.path
    )

    pct = [
        h
        for h in t.headers
        if any(
            k in nh(h)
            for k in [
                "percent",
                "percentage",
                "pct"
            ]
        )
    ]

    bad = []
    frac = []

    for h in pct:

        vv = [
            num(r[h])
            for r in t.rows
            if not missing(r[h])
        ]

        if any(
            v is None
            or v < 0
            or v > 100
            for v in vv
        ):
            bad.append(h)

        elif (
            vv
            and max(vv) <= 1
            and min(vv) >= 0
            and any(
                v not in {0, 1}
                for v in vv
            )
        ):
            frac.append(h)

    a.add(
        t.key,
        "percentage_range",
        "FAIL" if bad else "PASS",
        "Invalid percentage range"
        if bad
        else "Percentage columns within 0-100",
        bad or pct,
        "",
        t.path
    )

    if frac:

        a.add(
            t.key,
            "percentage_scale",
            "WARN",
            "Percentage-labelled columns look like 0-1 fractions",
            frac,
            "0-100 percentage points",
            t.path
        )


def unique(a, t, c, label):

    vv = [
        str(r[c]).strip()
        for r in t.rows
        if not missing(r[c])
    ]

    miss = sum(
        missing(r[c])
        for r in t.rows
    )

    dup = [
        k
        for k, v in Counter(vv).items()
        if v > 1
    ]

    a.add(
        t.key,
        "unique_" + label,
        "FAIL"
        if miss or dup
        else "PASS",
        f"{label} non-missing and unique"
        if not miss and not dup
        else f"{label} has missing/duplicates",
        f"missing={miss};duplicates={dup[:20]}",
        "",
        t.path
    )


# ============================================================
# S1 / S5 metadata validator
# ============================================================

def metadata(
    a,
    t,
    N,
    counts,
    cohort_label
):

    c = require(
        a,
        t,
        [
            "sample_id",
            "sampling_stage",
            "time_h"
        ]
    )

    if any(
        not v
        for v in c.values()
    ):
        return {}

    a.add(
        t.key,
        "row_count",
        "PASS"
        if t.n == N
        else "FAIL",
        "Row count",
        t.n,
        N,
        t.path
    )

    unique(
        a,
        t,
        c["sample_id"],
        "sample_id"
    )

    cc = Counter(
        stage(
            r[c["sampling_stage"]]
        )
        for r in t.rows
    )

    a.add(
        t.key,
        "stage_counts",
        "PASS"
        if dict(cc) == counts
        else "FAIL",
        "Stage counts",
        dict(cc),
        counts,
        t.path
    )

    err = []

    for i, r in enumerate(
        t.rows,
        2
    ):

        s = stage(
            r[c["sampling_stage"]]
        )

        v = num(
            r[c["time_h"]]
        )

        if s == "Baseline":

            if not missing(
                r[c["time_h"]]
            ):
                err.append(
                    f"row{i}:Baseline="
                    f"{r[c['time_h']]}"
                )

        else:

            ex = stage_time(s)

            if (
                ex is None
                or v is None
                or abs(v - ex) > 1e-8
            ):
                err.append(
                    f"row{i}:{s}:"
                    f"{r[c['time_h']]}"
                )

    a.add(
        t.key,
        "stage_time",
        "FAIL"
        if err
        else "PASS",
        "Stage/time consistency",
        err[:20],
        "Baseline=NA; injury stage numeric hour",
        t.path
    )

    co = col(t, "cohort")

    if co:

        vals = sorted({
            str(r[co]).strip()
            for r in t.rows
            if not missing(r[co])
        })

        a.add(
            t.key,
            "cohort_label",
            "PASS"
            if vals == [cohort_label]
            else "FAIL",
            "Cohort label",
            vals,
            cohort_label,
            t.path
        )

    else:

        a.add(
            t.key,
            "cohort_context",
            "PASS",
            "Cohort is defined by the dedicated table identity",
            cohort_label,
            cohort_label,
            t.path
        )

    cd = optional(
        a,
        t,
        "condition"
    )

    if cd:

        e = []

        for i, r in enumerate(
            t.rows,
            2
        ):

            s = stage(
                r[c["sampling_stage"]]
            )

            ex = (
                "Uninjured"
                if s == "Baseline"
                else "Injury"
            )

            if cond(r[cd]) != ex:

                e.append(
                    f"row{i}:{r[cd]}"
                )

        a.add(
            t.key,
            "condition",
            "FAIL"
            if e
            else "PASS",
            "Condition consistent with stage",
            e[:20],
            "Baseline=Uninjured; others=Injury",
            t.path
        )

    return {
        str(r[c["sample_id"]]).strip():
        stage(
            r[c["sampling_stage"]]
        )
        for r in t.rows
        if not missing(
            r[c["sample_id"]]
        )
    }


# ============================================================
# S2
# ============================================================

def s2(a, t, s1):

    c = require(
        a,
        t,
        [
            "sample_id",
            "sampling_stage"
        ]
    )

    if any(
        not v
        for v in c.values()
    ):
        return

    a.add(
        "S2",
        "row_count",
        "PASS"
        if t.n == 77
        else "FAIL",
        "S2 row count",
        t.n,
        77,
        t.path
    )

    unique(
        a,
        t,
        c["sample_id"],
        "sample_id"
    )

    ids = {
        str(r[c["sample_id"]]).strip()
        for r in t.rows
        if not missing(
            r[c["sample_id"]]
        )
    }

    if s1:

        a.add(
            "S2",
            "sample_set_matches_s1",
            "PASS"
            if ids == set(s1)
            else "FAIL",
            "S2 sample set vs S1",
            (
                f"missing="
                f"{sorted(set(s1) - ids)};"
                f"extra="
                f"{sorted(ids - set(s1))}"
            ),
            "exact S1 set",
            t.path
        )

    cc = Counter(
        stage(
            r[c["sampling_stage"]]
        )
        for r in t.rows
    )

    a.add(
        "S2",
        "stage_counts",
        "PASS"
        if dict(cc) == S1_COUNTS
        else "FAIL",
        "S2 stage counts",
        dict(cc),
        S1_COUNTS,
        t.path
    )

    QC_ALIASES = {

        "q20_percent": [
            "q20_percent",
            "q20",
            "q20_pct",
            "q20_clean_percent"
        ],

        "q30_percent": [
            "q30_percent",
            "q30",
            "q30_pct",
            "q30_clean_percent"
        ],

        "hisat2_alignment_percent": [
            "hisat2_alignment_percent",
            "overall_alignment_rate",
            "alignment_percent",
            "mapping_rate",
            "overall_alignment",
            "hisat2_overall_alignment_rate_percent"
        ],

        "featurecounts_assignment_percent": [
            "featurecounts_assignment_percent",
            "assignment_percent",
            "assigned_percent",
            "featurecounts_percent",
            "featurecounts_assignment_rate_percent"
        ]
    }

    m = {
        nh(h): h
        for h in t.headers
    }

    for name, aliases in QC_ALIASES.items():

        h = next(
            (
                m[nh(x)]
                for x in aliases
                if nh(x) in m
            ),
            None
        )

        if not h:

            a.add(
                "S2",
                "qc_" + name,
                "WARN",
                f"Expected QC column not found: {name}",
                path=t.path
            )

            continue

        vv = [
            num(r[h])
            for r in t.rows
        ]

        invalid = any(
            v is None
            or v < 0
            or v > 100
            for v in vv
        )

        observed = "invalid"

        if all(
            v is not None
            for v in vv
        ):
            observed = (
                f"{min(vv):.3f}-"
                f"{max(vv):.3f}"
            )

        a.add(
            "S2",
            "qc_" + name,
            "FAIL"
            if invalid
            else "PASS",
            f"{name} numeric/range",
            observed,
            "0-100",
            t.path
        )


# ============================================================
# S3 / S6 DEG summary
# ============================================================

def deg(
    a,
    t,
    expected,
    tested,
    baseN
):

    c = require(
        a,
        t,
        [
            "n_upregulated",
            "n_downregulated",
            "n_total_deg"
        ]
    )

    if any(
        not v
        for v in c.values()
    ):
        return

    sc = col(
        t,
        "sampling_stage"
    )

    cc = col(
        t,
        "contrast"
    )

    if not sc and not cc:

        a.add(
            t.key,
            "stage_identifier",
            "FAIL",
            "Need sampling_stage or contrast",
            path=t.path
        )

        return

    obs = {}
    arith = []

    for i, r in enumerate(
        t.rows,
        2
    ):

        if sc:

            s = stage(
                r[sc]
            )

        else:

            first = re.split(
                r"(?i)_?vs_?|versus",
                str(r[cc])
            )[0].strip("_ -")

            s = stage(first)

        u = integer(
            r[c["n_upregulated"]]
        )

        d = integer(
            r[c["n_downregulated"]]
        )

        z = integer(
            r[c["n_total_deg"]]
        )

        if s in obs:

            a.add(
                t.key,
                "duplicate_stage",
                "FAIL",
                "Duplicate DEG stage",
                s,
                "",
                t.path
            )

        obs[s] = (
            u,
            d,
            z
        )

        if (
            None in (u, d, z)
            or u + d != z
        ):

            arith.append(
                f"{s}:{u}+{d}!={z}"
            )

    a.add(
        t.key,
        "row_count",
        "PASS"
        if t.n == len(expected)
        else "FAIL",
        "DEG summary row count",
        t.n,
        len(expected),
        t.path
    )

    a.add(
        t.key,
        "deg_arithmetic",
        "FAIL"
        if arith
        else "PASS",
        "Up + Down = Total",
        arith[:20],
        "",
        t.path
    )

    bad = [
        (
            f"{s}:obs={obs.get(s)},"
            f"exp={v}"
        )
        for s, v in expected.items()
        if obs.get(s) != v
    ]

    a.add(
        t.key,
        "frozen_deg_counts",
        "FAIL"
        if bad
        else "PASS",
        "Frozen DEG counts",
        bad
        if bad
        else "all matched",
        "",
        t.path
    )

    ng = col(
        t,
        "n_genes_tested"
    )

    if ng:

        vv = [
            integer(r[ng])
            for r in t.rows
        ]

        a.add(
            t.key,
            "n_genes_tested",
            "PASS"
            if all(
                v == tested
                for v in vv
            )
            else "FAIL",
            "Gene universe",
            sorted(set(vv)),
            tested,
            t.path
        )

    nb = col(
        t,
        "n_samples_baseline"
    )

    if nb:

        vv = [
            integer(r[nb])
            for r in t.rows
        ]

        a.add(
            t.key,
            "baseline_n",
            "PASS"
            if all(
                v == baseN
                for v in vv
            )
            else "FAIL",
            "Baseline n",
            sorted(set(vv)),
            baseN,
            t.path
        )


# ============================================================
# S4 temporal annotation
# ============================================================

def s4(a, t):

    g = col(t, "gene_id")

    if not g:
        a.add(
            "S4",
            "gene_id",
            "FAIL",
            "gene_id column required",
            t.headers,
            "",
            t.path
        )
        return

    unique(a, t, g, "gene_id")

    # S4 is expected to contain the complete 18,364-gene
    # primary filtered universe.
    a.add(
        "S4",
        "row_count",
        "PASS" if t.n == 18364 else "FAIL",
        "Complete temporal-annotation gene universe",
        t.n,
        18364,
        t.path
    )

    # --------------------------------------------------------
    # Resolve actual frozen S4 columns
    # --------------------------------------------------------

    def exact_or_alias(*names):
        m = {nh(h): h for h in t.headers}
        for name in names:
            if nh(name) in m:
                return m[nh(name)]
        return None

    tau = exact_or_alias(
        "Tau_injury10",
        "tau_injury10",
        "tau"
    )

    spm = exact_or_alias(
        "Peak_SPM_injury10",
        "peak_spm_injury10",
        "peak_spm"
    )

    lrt = exact_or_alias(
        "LRT_dynamic_FDR005",
        "lrt_dynamic_fdr005"
    )

    high = exact_or_alias(
        "High_confidence_dynamic",
        "higher_amplitude_dynamic"
    )

    tsp = exact_or_alias(
        "Tau085_and_SPM050_injury10",
        "tau085_spm050_injury10"
    )

    tau_flag = exact_or_alias(
        "Tau_GE_0.85_injury10"
    )

    spm_flag = exact_or_alias(
        "Peak_SPM_GE_0.50_injury10"
    )

    pk = exact_or_alias(
        "Peak_stage_label_injury10",
        "Peak_stage_injury10",
        "observed_peak_stage"
    )

    required = {
        "Tau_injury10": tau,
        "Peak_SPM_injury10": spm,
        "LRT_dynamic_FDR005": lrt,
        "High_confidence_dynamic": high,
        "Tau085_and_SPM050_injury10": tsp,
        "Peak_stage_injury10": pk
    }

    miss = [k for k, v in required.items() if not v]

    a.add(
        "S4",
        "temporal_annotation_columns",
        "FAIL" if miss else "PASS",
        "Required temporal annotation columns",
        "missing=" + ",".join(miss) if miss else "all present",
        "all present",
        t.path
    )

    if miss:
        return

    # --------------------------------------------------------
    # Numeric ranges: Tau and SPM must lie within 0..1
    # --------------------------------------------------------

    for label, h in [
        ("Tau_injury10", tau),
        ("Peak_SPM_injury10", spm)
    ]:

        vv = [
            num(r[h])
            for r in t.rows
            if not missing(r[h])
        ]

        bad = any(
            v is None or v < 0 or v > 1
            for v in vv
        )

        a.add(
            "S4",
            label + "_range",
            "FAIL" if bad else "PASS",
            label + " values within 0-1",
            (
                f"n={len(vv)};"
                f"min={min(vv):.6g};"
                f"max={max(vv):.6g}"
                if vv else "n=0"
            ),
            "0-1",
            t.path
        )

    # --------------------------------------------------------
    # Boolean parsing
    # --------------------------------------------------------

    def bool_vector(h):
        return [boolean(r[h]) for r in t.rows]

    lrt_v = bool_vector(lrt)
    high_v = bool_vector(high)
    tsp_v = bool_vector(tsp)

    bad_bool = {
        "LRT_dynamic_FDR005": sum(x is None for x in lrt_v),
        "High_confidence_dynamic": sum(x is None for x in high_v),
        "Tau085_and_SPM050_injury10": sum(x is None for x in tsp_v)
    }

    a.add(
        "S4",
        "boolean_fields",
        "FAIL" if any(bad_bool.values()) else "PASS",
        "Boolean annotation fields are parseable",
        bad_bool,
        "0 unparsable",
        t.path
    )

    if any(bad_bool.values()):
        return

    # --------------------------------------------------------
    # Frozen counts
    # --------------------------------------------------------

    n_lrt = sum(lrt_v)
    n_high = sum(high_v)

    final_mask = [
        h and q
        for h, q in zip(high_v, tsp_v)
    ]

    n_final = sum(final_mask)

    a.add(
        "S4",
        "lrt_dynamic_count",
        "PASS" if n_lrt == 15322 else "FAIL",
        "Genes with LRT_dynamic_FDR005=TRUE",
        n_lrt,
        15322,
        t.path
    )

    a.add(
        "S4",
        "high_confidence_dynamic_count",
        "PASS" if n_high == 9539 else "FAIL",
        "Genes with High_confidence_dynamic=TRUE",
        n_high,
        9539,
        t.path
    )

    a.add(
        "S4",
        "final_temporal_specificity_count",
        "PASS" if n_final == 476 else "FAIL",
        (
            "High_confidence_dynamic AND "
            "Tau085_and_SPM050_injury10"
        ),
        n_final,
        476,
        t.path
    )

    # --------------------------------------------------------
    # Verify boolean threshold flags against numeric Tau/SPM
    # --------------------------------------------------------

    if tau_flag:

        bad = []

        for i, r in enumerate(t.rows, 2):
            tv = num(r[tau])
            bv = boolean(r[tau_flag])

            if bv is None:
                bad.append(i)
                continue

            # Undefined Tau is valid only when the threshold flag is FALSE.
            if tv is None:
                if bv:
                    bad.append(i)
                continue

            if bv != (tv >= 0.85):
                bad.append(i)

        a.add(
            "S4",
            "tau_threshold_flag",
            "FAIL" if bad else "PASS",
            "Tau_GE_0.85_injury10 agrees with numeric Tau",
            bad[:20],
            "Tau >= 0.85",
            t.path
        )

    if spm_flag:

        bad = []

        for i, r in enumerate(t.rows, 2):
            sv = num(r[spm])
            bv = boolean(r[spm_flag])

            if bv is None:
                bad.append(i)
                continue

            # Undefined SPM is valid only when the threshold flag is FALSE.
            if sv is None:
                if bv:
                    bad.append(i)
                continue

            if bv != (sv >= 0.50):
                bad.append(i)

        a.add(
            "S4",
            "spm_threshold_flag",
            "FAIL" if bad else "PASS",
            "Peak_SPM_GE_0.50_injury10 agrees with numeric SPM",
            bad[:20],
            "Peak SPM >= 0.50",
            t.path
        )

    # Verify combined Tau/SPM Boolean directly.
    bad = []

    for i, r in enumerate(t.rows, 2):

        tv = num(r[tau])
        sv = num(r[spm])
        bv = boolean(r[tsp])

        if bv is None:
            bad.append(i)
            continue

        # If Tau/SPM is undefined, the combined threshold flag must be FALSE.
        if tv is None or sv is None:
            if bv:
                bad.append(i)
            continue

        expected = (
            tv >= 0.85
            and sv >= 0.50
        )

        if bv != expected:
            bad.append(i)

    a.add(
        "S4",
        "combined_tau_spm_flag",
        "FAIL" if bad else "PASS",
        (
            "Tau085_and_SPM050_injury10 agrees "
            "with numeric thresholds"
        ),
        bad[:20],
        "Tau>=0.85 AND peak SPM>=0.50",
        t.path
    )

    # --------------------------------------------------------
    # Frozen 476-gene peak-stage distribution
    # --------------------------------------------------------

    cc = Counter(
        stage(r[pk])
        for r, keep in zip(t.rows, final_mask)
        if keep
    )

    a.add(
        "S4",
        "peak_distribution_476",
        "PASS" if dict(cc) == PEAK476 else "FAIL",
        "Observed peak-stage distribution among final 476 genes",
        dict(cc),
        PEAK476,
        t.path
    )

    # --------------------------------------------------------
    # Median TPM columns expected for all 11 stages
    # --------------------------------------------------------

    expected_tpm = [
        "Baseline",
        "R0h",
        "R1h",
        "R2h",
        "R6h",
        "R10h",
        "R14h",
        "R18h",
        "R36h",
        "R60h",
        "R72h"
    ]

    missing_tpm = [
        x for x in expected_tpm
        if nh(x) not in {nh(h) for h in t.headers}
    ]

    a.add(
        "S4",
        "median_tpm_columns",
        "FAIL" if missing_tpm else "PASS",
        "Stage-median TPM columns",
        missing_tpm if missing_tpm else "all 11 present",
        "Baseline + 10 injury stages",
        t.path
    )


# ============================================================
# S7 cross-cohort table
# ============================================================

def s7(a, t):

    c = require(
        a,
        t,
        [
            "primary_stage",
            "historical_stage",
            "spearman_rho"
        ]
    )

    if any(
        not v
        for v in c.values()
    ):
        return

    pairs = [
        (
            stage(
                r[c["primary_stage"]]
            ),
            stage(
                r[c["historical_stage"]]
            )
        )
        for r in t.rows
    ]

    exp = {
        (
            p,
            h
        )
        for p in PRIMARY
        for h in HIST
    }

    a.add(
        "S7",
        "row_count",
        "PASS"
        if t.n == 70
        else "FAIL",
        "S7 rows",
        t.n,
        70,
        t.path
    )

    a.add(
        "S7",
        "unique_pairs",
        "PASS"
        if len(set(pairs)) == 70
        else "FAIL",
        "Unique stage pairs",
        len(set(pairs)),
        70,
        t.path
    )

    a.add(
        "S7",
        "full_grid",
        "PASS"
        if set(pairs) == exp
        else "FAIL",
        "Complete 10x7 grid",
        (
            f"missing="
            f"{sorted(exp-set(pairs), key=str)};"
            f"extra="
            f"{sorted(set(pairs)-exp, key=str)}"
        ),
        "70 expected pairs",
        t.path
    )

    rr = [
        num(
            r[c["spearman_rho"]]
        )
        for r in t.rows
    ]

    invalid = any(
        v is None
        or v < -1
        or v > 1
        for v in rr
    )

    a.add(
        "S7",
        "spearman_range",
        "FAIL"
        if invalid
        else "PASS",
        "Spearman rho range",
        "",
        "-1..1",
        t.path
    )

    if all(
        v is not None
        for v in rr
    ):

        mean = statistics.mean(rr)
        med = statistics.median(rr)
        mn = min(rr)
        mx = max(rr)

        ok = (
            abs(mean - 0.37198)
            <= 0.002
            and abs(med - 0.42311)
            <= 0.002
            and abs(mn + 0.11664)
            <= 0.002
            and abs(mx - 0.72230)
            <= 0.002
        )

        a.add(
            "S7",
            "frozen_spearman_summary",
            "PASS"
            if ok
            else "FAIL",
            "Frozen Spearman summary",
            (
                f"mean={mean:.5f};"
                f"median={med:.5f};"
                f"min={mn:.5f};"
                f"max={mx:.5f}"
            ),
            (
                "mean=0.37198;"
                "median=0.42311;"
                "min=-0.11664;"
                "max=0.72230"
            ),
            t.path
        )

    ng = optional(
        a,
        t,
        "n_common_genes"
    )

    if ng:

        vv = [
            integer(r[ng])
            for r in t.rows
        ]

        a.add(
            "S7",
            "common_genes",
            "PASS"
            if all(
                v == 16918
                for v in vv
            )
            else "FAIL",
            "Common gene universe",
            sorted(set(vv)),
            16918,
            t.path
        )

    pr = optional(
        a,
        t,
        "pearson_r"
    )

    if pr:

        vv = [
            num(r[pr])
            for r in t.rows
        ]

        a.add(
            "S7",
            "pearson_range",
            "FAIL"
            if any(
                v is None
                or v < -1
                or v > 1
                for v in vv
            )
            else "PASS",
            "Pearson r range",
            "",
            "-1..1",
            t.path
        )

    sh = optional(
        a,
        t,
        "shared_deg_n"
    )

    dc = optional(
        a,
        t,
        "direction_concordance_percent"
    )

    sd = col(
        t,
        "shared_deg_same_direction_n"
    )

    if sh:

        vv = [
            integer(r[sh])
            for r in t.rows
        ]

        invalid = any(
            v is None
            or v < 0
            for v in vv
        )

        obs = "invalid"

        if all(
            v is not None
            for v in vv
        ):
            obs = (
                f"{min(vv)}-"
                f"{max(vv)}"
            )

        a.add(
            "S7",
            "shared_deg",
            "FAIL"
            if invalid
            else "PASS",
            "Shared DEG counts",
            obs,
            "non-negative",
            t.path
        )

    if dc:

        vv = [
            num(r[dc])
            for r in t.rows
        ]

        a.add(
            "S7",
            "direction_concordance",
            "FAIL"
            if any(
                v is None
                or v < 0
                or v > 100
                for v in vv
            )
            else "PASS",
            "Direction concordance",
            "",
            "0-100",
            t.path
        )

    if (
        sh
        and dc
        and sd
    ):

        bad = []

        for i, r in enumerate(
            t.rows,
            2
        ):

            n = integer(r[sh])
            same = integer(r[sd])
            pct = num(r[dc])

            if (
                None in (
                    n,
                    same,
                    pct
                )
                or same > n
            ):

                bad.append(i)
                continue

            calc = (
                100 * same / n
                if n
                else 0
            )

            if abs(
                pct - calc
            ) > 0.15:

                bad.append(i)

        a.add(
            "S7",
            "direction_arithmetic",
            "FAIL"
            if bad
            else "PASS",
            "Direction concordance arithmetic",
            bad[:20],
            "",
            t.path
        )

    pt = col(
        t,
        "primary_time_h"
    )

    ht = col(
        t,
        "historical_time_h"
    )

    gap = col(
        t,
        "absolute_time_gap_h"
    )

    near = col(
        t,
        "nearest_time_pair"
    )

    gaps = {}
    bad = []

    for i, r in enumerate(
        t.rows,
        2
    ):

        ps = stage(
            r[c["primary_stage"]]
        )

        hs = stage(
            r[c["historical_stage"]]
        )

        x = (
            num(r[pt])
            if pt
            else stage_time(ps)
        )

        y = (
            num(r[ht])
            if ht
            else stage_time(hs)
        )

        if (
            x is None
            or y is None
        ):

            bad.append(i)
            continue

        g = abs(
            x - y
        )

        gaps[
            (
                ps,
                hs
            )
        ] = g

        if gap:

            gv = num(
                r[gap]
            )

            if (
                gv is None
                or abs(gv - g) > 1e-8
            ):
                bad.append(i)

    a.add(
        "S7",
        "time_gap",
        "FAIL"
        if bad
        else "PASS",
        "Stage/time-gap consistency",
        bad[:20],
        "",
        t.path
    )

    if (
        near
        and gaps
    ):

        ex = set()

        for p in PRIMARY:

            m = min(
                gaps[
                    (
                        p,
                        h
                    )
                ]
                for h in HIST
            )

            ex |= {
                (
                    p,
                    h
                )
                for h in HIST
                if abs(
                    gaps[
                        (
                            p,
                            h
                        )
                    ] - m
                ) < 1e-8
            }

        ob = set()
        badb = []

        for i, r in enumerate(
            t.rows,
            2
        ):

            b = boolean(
                r[near]
            )

            pair = (
                stage(
                    r[c["primary_stage"]]
                ),
                stage(
                    r[c["historical_stage"]]
                )
            )

            if b is None:
                badb.append(i)

            elif b:
                ob.add(pair)

        ok = (
            not badb
            and ob == ex
        )

        a.add(
            "S7",
            "nearest_pairs",
            "PASS"
            if ok
            else "FAIL",
            "Nearest-time flags with ties",
            (
                f"bad_bool={badb};"
                f"observed="
                f"{sorted(ob, key=str)}"
            ),
            sorted(
                ex,
                key=str
            ),
            t.path
        )


# ============================================================
# S8 robustness table
# ============================================================

def s8(a, t):

    sc = col(t, "scenario_id")

    # Current S8 uses Analysis_type +
    # Scenario_or_removed_sample as a composite scenario key.
    m = {nh(h): h for h in t.headers}

    at = m.get(nh("Analysis_type"))

    if not sc or not at:

        a.add(
            "S8",
            "scenario_columns",
            "FAIL",
            (
                "Analysis_type and "
                "Scenario_or_removed_sample are required"
            ),
            t.headers,
            "",
            t.path
        )
        return

    keys = []
    missing_key = []

    for i, r in enumerate(t.rows, 2):

        x = str(r[at]).strip() if not missing(r[at]) else ""
        y = str(r[sc]).strip() if not missing(r[sc]) else ""

        if not x or not y:
            missing_key.append(i)

        keys.append(x + "::" + y)

    dup = [
        k for k, v in Counter(keys).items()
        if v > 1
    ]

    a.add(
        "S8",
        "composite_scenario_id",
        "FAIL" if missing_key or dup else "PASS",
        (
            "Analysis_type + Scenario_or_removed_sample "
            "uniquely identifies each row"
        ),
        (
            f"missing_rows={missing_key};"
            f"duplicates={dup}"
        ),
        "unique composite keys",
        t.path
    )

    nr = optional(
        a,
        t,
        "n_samples_remaining"
    )

    ne = col(
        t,
        "n_excluded"
    )

    r2 = optional(
        a,
        t,
        "permanova_r2"
    )

    pp = optional(
        a,
        t,
        "permanova_p"
    )

    dp = optional(
        a,
        t,
        "permdisp_p"
    )

    mc = optional(
        a,
        t,
        "matrix_correlation_vs_primary"
    )

    mae = optional(
        a,
        t,
        "matrix_mean_absolute_difference"
    )

    if nr:

        vv = [
            integer(r[nr])
            for r in t.rows
        ]

        a.add(
            "S8",
            "remaining_n",
            "FAIL"
            if any(
                v is None
                or v < 1
                or v > 77
                for v in vv
            )
            else "PASS",
            "Samples remaining",
            "",
            "1-77",
            t.path
        )

    if (
        nr
        and ne
    ):

        bad = [
            i
            for i, r in enumerate(
                t.rows,
                2
            )
            if (
                integer(r[nr]) is None
                or integer(r[ne]) is None
                or integer(r[nr])
                + integer(r[ne])
                != 77
            )
        ]

        a.add(
            "S8",
            "sample_arithmetic",
            "FAIL"
            if bad
            else "PASS",
            "n_remaining + n_excluded = 77",
            bad[:20],
            "",
            t.path
        )

    for name, h, lo, hi in [

        (
            "permanova_r2",
            r2,
            0,
            1
        ),

        (
            "permanova_p",
            pp,
            0,
            1
        ),

        (
            "permdisp_p",
            dp,
            0,
            1
        ),

        (
            "matrix_correlation",
            mc,
            -1,
            1
        )
    ]:

        if not h:
            continue

        vv = [
            num(r[h])
            for r in t.rows
            if not missing(r[h])
        ]

        a.add(
            "S8",
            name + "_range",
            "FAIL"
            if any(
                v is None
                or v < lo
                or v > hi
                for v in vv
            )
            else "PASS",
            name + " range",
            f"n={len(vv)}",
            f"{lo}..{hi}",
            t.path
        )

    if mae:

        vv = [
            num(r[mae])
            for r in t.rows
            if not missing(r[mae])
        ]

        a.add(
            "S8",
            "mae_range",
            "FAIL"
            if any(
                v is None
                or v < 0
                for v in vv
            )
            else "PASS",
            "Matrix MAE non-negative",
            f"n={len(vv)}",
            ">=0",
            t.path
        )


# ============================================================
# Output
# ============================================================

def write(a, res, out):

    out.mkdir(
        parents=True,
        exist_ok=True
    )

    rep = (
        out
        / "16A_supplementary_tables_validation_report.tsv"
    )

    with rep.open(
        "w",
        encoding="utf-8",
        newline=""
    ) as f:

        w = csv.writer(
            f,
            delimiter="\t",
            lineterminator="\n"
        )

        w.writerow([
            "table",
            "check_id",
            "status",
            "message",
            "observed",
            "expected",
            "file"
        ])

        w.writerows(
            a.rows
        )

    by = defaultdict(list)

    for r in a.rows:
        by[r[0]].append(r)

    summ = (
        out
        / "16A_supplementary_tables_validation_summary.tsv"
    )

    with summ.open(
        "w",
        encoding="utf-8",
        newline=""
    ) as f:

        w = csv.writer(
            f,
            delimiter="\t",
            lineterminator="\n"
        )

        w.writerow([
            "table",
            "status",
            "pass",
            "warn",
            "fail",
            "file"
        ])

        for k in [
            f"S{i}"
            for i in range(1, 9)
        ]:

            x = by[k]

            P = sum(
                r[2] == "PASS"
                for r in x
            )

            W = sum(
                r[2] == "WARN"
                for r in x
            )

            F = sum(
                r[2] == "FAIL"
                for r in x
            )

            status = (
                "FAIL"
                if F
                else (
                    "WARN"
                    if W
                    else "PASS"
                )
            )

            w.writerow([
                k,
                status,
                P,
                W,
                F,
                res.get(k) or ""
            ])

    rf = (
        out
        / "16A_resolved_table_files.tsv"
    )

    with rf.open(
        "w",
        encoding="utf-8",
        newline=""
    ) as f:

        w = csv.writer(
            f,
            delimiter="\t",
            lineterminator="\n"
        )

        w.writerow([
            "table",
            "file"
        ])

        for k in [
            f"S{i}"
            for i in range(1, 9)
        ]:

            w.writerow([
                k,
                res.get(k) or ""
            ])

    P = sum(
        r[2] == "PASS"
        for r in a.rows
    )

    W = sum(
        r[2] == "WARN"
        for r in a.rows
    )

    F = sum(
        r[2] == "FAIL"
        for r in a.rows
    )

    overall = (
        "FAIL"
        if F
        else "PASS"
    )

    log = (
        out
        / "16A_supplementary_tables_validation.log"
    )

    with log.open(
        "w",
        encoding="utf-8"
    ) as f:

        f.write(
            "=" * 72
            + "\n"
        )

        f.write(
            "SCIENTIFIC DATA S1-S8 PRE-SUBMISSION AUDIT\n"
        )

        f.write(
            "=" * 72
            + "\n"
        )

        f.write(
            f"OVERALL: {overall}\n"
        )

        f.write(
            f"PASS={P} "
            f"WARN={W} "
            f"FAIL={F}\n"
        )

        for k in [
            f"S{i}"
            for i in range(1, 9)
        ]:

            x = by[k]

            p = sum(
                r[2] == "PASS"
                for r in x
            )

            w = sum(
                r[2] == "WARN"
                for r in x
            )

            q = sum(
                r[2] == "FAIL"
                for r in x
            )

            st = (
                "FAIL"
                if q
                else (
                    "WARN"
                    if w
                    else "PASS"
                )
            )

            f.write(
                f"{k}: {st} "
                f"PASS={p} "
                f"WARN={w} "
                f"FAIL={q}\n"
            )

            for r in x:

                if r[2] in {
                    "WARN",
                    "FAIL"
                }:

                    f.write(
                        f"  [{r[2]}] "
                        f"{r[1]}: "
                        f"{r[3]} | "
                        f"observed={r[4]} | "
                        f"expected={r[5]}\n"
                    )

    print(
        "=" * 72
    )

    print(
        "SCIENTIFIC DATA S1-S8 PRE-SUBMISSION AUDIT"
    )

    print(
        "=" * 72
    )

    print(
        f"OVERALL: {overall}   "
        f"PASS={P} "
        f"WARN={W} "
        f"FAIL={F}"
    )

    for k in [
        f"S{i}"
        for i in range(1, 9)
    ]:

        x = by[k]

        p = sum(
            r[2] == "PASS"
            for r in x
        )

        w = sum(
            r[2] == "WARN"
            for r in x
        )

        q = sum(
            r[2] == "FAIL"
            for r in x
        )

        st = (
            "FAIL"
            if q
            else (
                "WARN"
                if w
                else "PASS"
            )
        )

        print(
            f"{k}: "
            f"{st:<4} "
            f"PASS={p:<3} "
            f"WARN={w:<3} "
            f"FAIL={q:<3}"
        )

    print(
        "-" * 72
    )

    print(
        "Detailed:",
        rep
    )

    print(
        "Summary :",
        summ
    )

    print(
        "Log     :",
        log
    )

    print(
        "=" * 72
    )


# ============================================================
# Main
# ============================================================

def main():

    ap = argparse.ArgumentParser(
        description=(
            "Pre-submission audit of "
            "Scientific Data Supplementary Tables S1-S8"
        )
    )

    ap.add_argument(
        "--input-dir",
        default=None,
        help=(
            "Directory containing S1-S8. "
            "Default: "
            "submission_package_scientific_data/"
            "03_Supplementary_Tables "
            "if present; otherwise current directory."
        )
    )

    ap.add_argument(
        "--output-dir",
        default="16A_supplementary_table_validation"
    )

    ap.add_argument(
        "--strict",
        action="store_true",
        help=(
            "Convert WARN checks to FAIL."
        )
    )

    for i in range(
        1,
        9
    ):

        ap.add_argument(
            f"--s{i}",
            default=None,
            help=(
                f"Explicit path to Supplementary "
                f"Table S{i}; overrides discovery."
            )
        )

    z = ap.parse_args()

    preferred = Path(
        "submission_package_scientific_data/"
        "03_Supplementary_Tables"
    )

    if z.input_dir:

        root = Path(
            z.input_dir
        )

    elif preferred.is_dir():

        root = preferred

    else:

        root = Path(".")

    a = Audit(
        z.strict
    )

    res = {}

    if not root.exists():

        print(
            "ERROR input dir not found:",
            root,
            file=sys.stderr
        )

        return 1

    # --------------------------------------------------------
    # Resolve S1-S8 files
    # --------------------------------------------------------

    for i in range(
        1,
        9
    ):

        k = f"S{i}"

        ov = getattr(
            z,
            f"s{i}"
        )

        if ov:

            p = Path(ov)

            if p.is_file():

                res[k] = p

                a.add(
                    k,
                    "file_discovery",
                    "PASS",
                    "Explicit file",
                    p,
                    "",
                    p
                )

            else:

                res[k] = None

                a.add(
                    k,
                    "file_discovery",
                    "FAIL",
                    "Explicit file not found",
                    p,
                    "",
                    p
                )

        else:

            p, cand = discover(
                root,
                i
            )

            res[k] = p

            if p:

                a.add(
                    k,
                    "file_discovery",
                    "PASS",
                    "Auto-discovered file",
                    p,
                    "",
                    p
                )

            else:

                a.add(
                    k,
                    "file_discovery",
                    "FAIL",
                    (
                        "No unique file found; "
                        "use --sN"
                    ),
                    cand,
                    "",
                    root
                )

    # --------------------------------------------------------
    # Read files
    # --------------------------------------------------------

    tables = {}

    for k, p in res.items():

        if not p:
            continue

        try:

            t = read_table(
                k,
                p
            )

            tables[k] = t

            a.add(
                k,
                "file_read",
                "PASS",
                "Read table",
                (
                    f"{t.n} rows x "
                    f"{len(t.headers)} cols"
                ),
                "",
                p
            )

            generic(
                a,
                t
            )

        except Exception as e:

            a.add(
                k,
                "file_read",
                "FAIL",
                f"Read failed: {e}",
                "",
                "",
                p
            )

    # --------------------------------------------------------
    # Table-specific validation
    # --------------------------------------------------------

    if "S1" in tables:

        m1 = metadata(
            a,
            tables["S1"],
            77,
            S1_COUNTS,
            "2026_primary"
        )

    else:

        m1 = {}

    if "S2" in tables:

        s2(
            a,
            tables["S2"],
            m1
        )

    if "S3" in tables:

        deg(
            a,
            tables["S3"],
            DEG2026,
            18364,
            7
        )

    if "S4" in tables:

        s4(
            a,
            tables["S4"]
        )

    if "S5" in tables:

        metadata(
            a,
            tables["S5"],
            60,
            S5_COUNTS,
            "2021_historical_reference"
        )

    if "S6" in tables:

        deg(
            a,
            tables["S6"],
            DEG2021,
            17328,
            6
        )

    if "S7" in tables:

        s7(
            a,
            tables["S7"]
        )

    if "S8" in tables:

        s8(
            a,
            tables["S8"]
        )

    # --------------------------------------------------------
    # Write reports
    # --------------------------------------------------------

    write(
        a,
        res,
        Path(
            z.output_dir
        )
    )

    return (
        1
        if a.failed()
        else 0
    )


if __name__ == "__main__":
    sys.exit(
        main()
    )
